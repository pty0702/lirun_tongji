# -*- coding: utf-8 -*-
"""
利润统计核心逻辑模块

职责：
- 字段识别
- 数据清洗
- 成本表处理
- 订单表处理
- 利润计算
- Excel 输出
- 命令行测试入口
"""

import re
import os
from datetime import datetime, timedelta

import pandas as pd
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, numbers
from openpyxl.utils import get_column_letter


# ============================================================
# 常量
# ============================================================

INVALID_CODES = {'-', '--', '无', 'nan', 'NaN', 'None', '', 'NONE', 'none', 'NULL', 'null'}

ORDER_COLUMN_ALIASES = {
    '商品ID': ['商品ID', '商品id', '产品ID', '产品id', '货品ID', '货品id', '商品编号', '商品编码'],
    '商家编码': ['商家编码', '商品编码', 'SKU编码', 'SKU', 'sku', '规格编码', '商品条码'],
    '商品数量': ['商品数量', '数量', '购买数量', '件数', '下单数量'],
    '订单提交时间': ['订单提交时间', '提交时间', '下单时间', '创建时间', '订单时间', '订单创建时间'],
    '商家收入': ['商家收入金额', '商家收入', '商家实收', '结算金额', '实收金额', '收入金额', '收入'],
    '订单号': ['主订单编号', '子订单编号', '订单号', '订单编号', '主订单号', '子订单号'],
    '订单状态': ['订单状态', '订单状态名称', '状态', '交易状态', '售后状态', '订单状态描述'],
}

COST_CODE_ALIASES = [
    '商品编码', '商品ID', '商家编码', '编码', '名称', '商品名称', '品名',
    '产品编码', '产品ID', '货号', '产品名称', '货品编码', '货品名称',
    '商品条码', '规格编码', 'SKU', 'sku',
]

COST_PRICE_ALIASES = [
    '总计成本', '成本', '成本价', '成本单价', '单价', '价格', '金额', '成本金额',
    '总成本', '合计成本', '进货价', '进货成本', '采购成本', '采购价',
]


# ============================================================
# 字段识别
# ============================================================

def normalize(s):
    """去除前后空格"""
    if isinstance(s, str):
        return s.strip()
    return s


def find_column(df, aliases):
    """在 DataFrame 中查找匹配别名的列名，返回第一个匹配的原始列名"""
    normalized_cols = {normalize(c): c for c in df.columns}
    for alias in aliases:
        for nc, oc in normalized_cols.items():
            if nc == alias or nc == normalize(alias):
                return oc
    # 模糊匹配：包含关键字
    for alias in aliases:
        for nc, oc in normalized_cols.items():
            if alias in nc:
                return oc
    return None


def identify_order_columns(df):
    """识别订单表中的关键字段，返回字段名映射"""
    result = {}
    for key, aliases in ORDER_COLUMN_ALIASES.items():
        col = find_column(df, aliases)
        if col:
            result[key] = col
    # 必需字段检查
    required = ['商品ID', '商家编码', '商品数量', '订单提交时间', '商家收入']
    missing = [r for r in required if r not in result]
    if missing:
        raise ValueError(f"订单表缺少必要字段: {', '.join(missing)}，请检查表头")
    return result


def identify_cost_columns(df):
    """识别成本表中的编码列和成本列"""
    code_col = find_column(df, COST_CODE_ALIASES)
    if code_col is None:
        raise ValueError("无法识别成本表编码列，请检查成本表表头")

    cost_col = find_column(df, COST_PRICE_ALIASES)
    if cost_col is None:
        raise ValueError("无法识别成本表成本列，请检查成本表表头")

    return code_col, cost_col


def get_unique_statuses(order_filepath):
    """读取订单表中的唯一订单状态列表"""
    try:
        df = pd.read_excel(order_filepath)
        col_map = identify_order_columns(df)
        status_col = col_map.get('订单状态', None)
        if status_col is None:
            return []
        statuses = df[status_col].dropna().unique()
        return sorted([clean_str(s) for s in statuses if clean_str(s)])
    except Exception:
        return []


# ============================================================
# 数据清洗
# ============================================================

def clean_str(val):
    """清洗字符串值"""
    if val is None:
        return ''
    if isinstance(val, (int, float)):
        if pd.isna(val):
            return ''
        # 大整数转字符串，避免科学计数法
        if isinstance(val, float) and val == int(val):
            return str(int(val))
        return str(val)
    s = str(val).strip()
    if s in INVALID_CODES:
        return ''
    # 去除制表符、换行等
    s = s.replace('\t', '').replace('\n', '').replace('\r', '')
    return s.strip()


def clean_numeric(val):
    """清洗数值，返回 float 或 NaN"""
    if val is None:
        return np.nan
    if isinstance(val, (int, float)):
        return float(val) if not pd.isna(val) else np.nan
    s = str(val).strip()
    if s in INVALID_CODES or s == '':
        return np.nan
    # 去除逗号、人民币符号、空格
    s = s.replace(',', '').replace('，', '').replace('￥', '').replace('¥', '').replace(' ', '')
    # 去除制表符
    s = s.replace('\t', '').replace('\n', '').replace('\r', '')
    try:
        return float(s)
    except ValueError:
        return np.nan


def clean_datetime(val):
    """清洗日期时间值，返回 datetime 或 NaT"""
    if val is None:
        return pd.NaT
    if isinstance(val, datetime):
        return val
    if isinstance(val, pd.Timestamp):
        return val.to_pydatetime()
    s = str(val).strip().replace('\t', '').replace('\n', '').replace('\r', '')
    if s in INVALID_CODES or s == '':
        return pd.NaT
    formats = [
        '%Y-%m-%d %H:%M:%S',
        '%Y-%m-%d %H:%M',
        '%Y/%m/%d %H:%M:%S',
        '%Y/%m/%d %H:%M',
        '%Y-%m-%d',
        '%Y/%m/%d',
        '%m/%d/%Y %H:%M:%S',
        '%m/%d/%Y',
    ]
    for fmt in formats:
        try:
            return datetime.strptime(s, fmt)
        except ValueError:
            continue
    # 最后尝试 pandas 解析
    try:
        return pd.to_datetime(s).to_pydatetime()
    except Exception:
        return pd.NaT


def is_invalid_code(val):
    """判断商家编码是否为无效值"""
    s = clean_str(val)
    return s == '' or s.lower() in {k.lower() for k in INVALID_CODES}


# ============================================================
# 成本表处理
# ============================================================

def process_cost_table(filepath):
    """读取并处理成本表，返回 (cost_map, cost_snapshot_records)

    cost_map: {商家编码: 成本单价}  仅包含正常可用的编码
    cost_snapshot_records: 成本表快照明细列表
    """
    df = pd.read_excel(filepath)
    code_col, cost_col = identify_cost_columns(df)

    records = []
    # 第一步：清洗数据并分组
    code_groups = {}  # code -> [(cost_value, row_idx)]

    for idx, row in df.iterrows():
        raw_code = row[code_col]
        raw_cost = row[cost_col]

        code = clean_str(raw_code)
        cost = clean_numeric(raw_cost)

        # 判断状态
        if code == '':
            status = '编码为空'
            reason = '成本表编码为空'
        elif pd.isna(cost):
            status = '成本异常'
            reason = '成本金额为空或无法转换为数字'
        else:
            status = '正常'
            reason = ''

        records.append({
            '商家编码': code if code else str(raw_code).strip(),
            '成本单价': cost if not pd.isna(cost) else None,
            '状态': status,
            '原因': reason,
            '_raw_code': code,
            '_raw_cost': cost,
        })

        if status == '正常':
            if code not in code_groups:
                code_groups[code] = []
            code_groups[code].append(cost)

    # 第二步：检测成本重复（同一编码对应多个不同成本）
    cost_map = {}
    for code, costs in code_groups.items():
        unique_costs = list(set(costs))
        if len(unique_costs) > 1:
            # 成本重复
            for rec in records:
                if rec['_raw_code'] == code and rec['状态'] == '正常':
                    rec['状态'] = '成本重复'
                    rec['原因'] = f"该编码在成本表中存在多个不同成本: {unique_costs}"
        else:
            cost_map[code] = unique_costs[0]

    # 清理内部字段
    for rec in records:
        del rec['_raw_code']
        del rec['_raw_cost']

    return cost_map, records


# ============================================================
# 订单表处理与利润计算
# ============================================================

def _vectorized_clean(df, col_map):
    """向量化清洗订单表各列，返回清洗后的 DataFrame（含中间列 _product_id, _merchant_code, _qty, _income, _submit_time, _order_no, _status_val）

    比逐行 iterrows 快约 19 倍（91000行: 0.24s vs 4.61s）
    """
    id_col = col_map['商品ID']
    code_col = col_map['商家编码']
    qty_col = col_map['商品数量']
    time_col = col_map['订单提交时间']
    income_col = col_map['商家收入']
    order_col = col_map.get('订单号', None)
    status_col = col_map.get('订单状态', None)

    n = len(df)

    # 字符串列：转字符串 + 去前后空格
    df['_product_id'] = df[id_col].astype(str).str.strip()
    df['_merchant_code'] = df[code_col].astype(str).str.strip()
    df['_order_no'] = df[order_col].astype(str).str.strip() if order_col else pd.Series([''] * n, index=df.index)
    df['_status_val'] = df[status_col].astype(str).str.strip() if status_col else pd.Series([''] * n, index=df.index)

    # 清理无效字符串值（将 INVALID_CODES 替换为空字符串）
    invalid_lower = {str(k).lower() for k in INVALID_CODES}
    for col_name in ['_product_id', '_merchant_code', '_order_no', '_status_val']:
        s = df[col_name]
        mask = s.isin(INVALID_CODES) | s.str.lower().isin(invalid_lower) | s.isna() | (s == 'nan') | (s == 'NaN') | (s == 'None') | (s == '')
        # 对于已经是空或 nan 的不需要再处理
        # 将明显的无效值设为空字符串
        df.loc[mask, col_name] = ''
        # 去掉制表符、换行
        df[col_name] = df[col_name].str.replace('\t', '', regex=False).str.replace('\n', '', regex=False).str.replace('\r', '', regex=False)

    # 数值列
    df['_qty'] = pd.to_numeric(df[qty_col], errors='coerce')
    df['_income'] = pd.to_numeric(df[income_col], errors='coerce')

    # 日期列
    df['_submit_time'] = pd.to_datetime(df[time_col], errors='coerce')

    return df


def process_orders(order_filepath, cost_map, start_date=None, end_date=None, order_statuses=None, df=None, col_map=None):
    """读取订单表、清洗、匹配成本、计算利润

    参数:
        order_statuses: 订单状态列表（多选），为 None 或空列表时不过滤
        df: 可选，已读取的 DataFrame（避免重复 read_excel）
        col_map: 可选，已识别的字段映射（避免重复识别）

    返回:
        detail_rows: 商品ID编码明细节记录
        summary_rows: 商品ID汇总记录
        anomaly_rows: 异常明细记录
    """
    # 仅在未传入时读取和识别（避免重复读取，大文件可节省 ~38秒）
    if df is None:
        df = pd.read_excel(order_filepath)
    if col_map is None:
        col_map = identify_order_columns(df)

    id_col = col_map['商品ID']
    code_col = col_map['商家编码']
    qty_col = col_map['商品数量']
    time_col = col_map['订单提交时间']
    income_col = col_map['商家收入']
    order_col = col_map.get('订单号', None)
    status_col = col_map.get('订单状态', None)

    # 向量化清洗
    df = _vectorized_clean(df, col_map)

    # 构建筛选掩码
    valid_mask = pd.Series(True, index=df.index)

    # 日期筛选
    if start_date:
        valid_mask &= df['_submit_time'].notna() & (df['_submit_time'] >= start_date)
    if end_date:
        valid_mask &= df['_submit_time'].notna() & (df['_submit_time'] <= end_date)

    # 订单状态筛选
    if order_statuses and status_col:
        valid_mask &= df['_status_val'].isin(order_statuses)

    # 在筛选后的数据上做异常检测
    # 使用索引来追踪
    working_df = df[valid_mask].copy()
    # 保留原始行索引以便回溯
    working_df['_orig_idx'] = working_df.index

    # --- 异常检测（优先级：编码为空 > 数量异常 > 收入异常 > 成本缺失）---
    # 各异常掩码
    code_invalid_mask = (
        (working_df['_merchant_code'] == '') |
        working_df['_merchant_code'].str.lower().isin({str(k).lower() for k in INVALID_CODES})
    )
    qty_invalid_mask = working_df['_qty'].isna() | (working_df['_qty'] <= 0)
    income_invalid_mask = working_df['_income'].isna()
    cost_missing_mask = ~working_df['_merchant_code'].isin(cost_map)

    # 按优先级：已被更高优先级标记的行，不再归入低优先级
    is_anomaly = pd.Series(False, index=working_df.index)
    anomaly_type_col = pd.Series('', index=working_df.index)
    anomaly_reason_col = pd.Series('', index=working_df.index)

    # 1. 编码为空
    mask1 = code_invalid_mask & ~is_anomaly
    is_anomaly |= mask1
    anomaly_type_col[mask1] = '商家编码为空'
    anomaly_reason_col[mask1] = '商家编码为空或为无效值'

    # 2. 数量异常
    mask2 = qty_invalid_mask & ~is_anomaly
    is_anomaly |= mask2
    # 原因中包含原始值
    anomaly_reason_col[mask2] = working_df.loc[mask2, qty_col].apply(
        lambda v: f'商品数量为空、无法转数字或小于等于0: {v}'
    )

    # 3. 收入异常
    mask3 = income_invalid_mask & ~is_anomaly
    is_anomaly |= mask3
    anomaly_reason_col[mask3] = working_df.loc[mask3, income_col].apply(
        lambda v: f'商家收入为空或无法转数字: {v}'
    )

    # 4. 成本缺失
    mask4 = cost_missing_mask & ~is_anomaly
    is_anomaly |= mask4
    anomaly_type_col[mask4] = '成本缺失'
    anomaly_reason_col[mask4] = working_df.loc[mask4, '_merchant_code'].apply(
        lambda v: f'商家编码"{v}"在成本表中未找到'
    )

    normal_mask = ~is_anomaly

    # --- 构建异常明细 ---
    anomaly_rows = []
    if is_anomaly.any():
        anomaly_subset = working_df[is_anomaly]
        for _, row in anomaly_subset.iterrows():
            orig_idx = row['_orig_idx']
            orig_row = df.loc[orig_idx]

            anomaly_rows.append({
                '异常类型': anomaly_type_col[row.name],
                '商品ID': row['_product_id'],
                '商家编码': row['_merchant_code'],
                '商品数量': row['_qty'] if not pd.isna(row['_qty']) else orig_row[qty_col],
                '商家收入': row['_income'] if not pd.isna(row['_income']) else orig_row[income_col],
                '订单提交时间': row['_submit_time'] if not pd.isna(row['_submit_time']) else orig_row[time_col],
                '订单号': row['_order_no'],
                '原因': anomaly_reason_col[row.name],
            })

    # --- 构建正常数据 ---
    normal_rows = []
    if normal_mask.any():
        normal_subset = working_df[normal_mask]
        for _, row in normal_subset.iterrows():
            merchant_code = row['_merchant_code']
            if merchant_code not in cost_map:
                # 理论不应走到这里（已被成本缺失掩码覆盖），但保留安全检查
                continue
            normal_rows.append({
                '商品ID': row['_product_id'],
                '商家编码': merchant_code,
                '商品数量': row['_qty'],
                '商家收入': row['_income'],
                '订单提交时间': row['_submit_time'] if not pd.isna(row['_submit_time']) else None,
                '订单号': row['_order_no'],
                '成本单价': cost_map[merchant_code],
            })

    # ========== 汇总计算 ==========

    # --- 按 商品ID + 商家编码 汇总 ---
    detail_df = pd.DataFrame(normal_rows)
    if len(detail_df) > 0:
        grouped = detail_df.groupby(['商品ID', '商家编码'], dropna=False)

        detail_records = []
        for (pid, mcode), grp in grouped:
            order_lines = len(grp)
            total_qty = grp['商品数量'].sum()
            cost_price = grp['成本单价'].iloc[0]
            total_cost = cost_price * total_qty
            total_income = grp['商家收入'].sum()
            total_profit = total_income - total_cost

            detail_records.append({
                '商品ID': pid,
                '商家编码': mcode,
                '成本单价': cost_price,
                '编码对应订单行数': order_lines,
                '编码对应商品数量': total_qty,
                '编码对应总成本': round(total_cost, 2),
                '编码对应总收入': round(total_income, 2),
                '编码对应总利润': round(total_profit, 2),
                '_total_profit': total_profit,
            })

        # 计算 ID对应总利润
        id_profit_map = {}
        for rec in detail_records:
            pid = rec['商品ID']
            id_profit_map[pid] = id_profit_map.get(pid, 0) + rec['_total_profit']

        for rec in detail_records:
            rec['ID对应总利润'] = round(id_profit_map[rec['商品ID']], 2)
            rec['状态'] = '正常'
            del rec['_total_profit']
    else:
        detail_records = []

    # --- 按 商品ID 汇总 ---
    normal_ids = set()
    if len(detail_df) > 0:
        summary_groups = detail_df.groupby('商品ID', dropna=False)
        summary_records = []
        for pid, grp in summary_groups:
            normal_ids.add(pid)
            summary_records.append({
                '商品ID': pid,
                '订单行数': len(grp),
                '商品数量合计': grp['商品数量'].sum(),
                '总收入': round(grp['商家收入'].sum(), 2),
                '总成本': round((grp['成本单价'] * grp['商品数量']).sum(), 2),
                '总利润': round(grp['商家收入'].sum() - (grp['成本单价'] * grp['商品数量']).sum(), 2),
                '_temp_id': pid,
            })
    else:
        summary_records = []

    # 为 summary 添加异常统计
    anomaly_df = pd.DataFrame(anomaly_rows)
    if len(anomaly_df) > 0:
        anomaly_summary = {}
        for _, arow in anomaly_df.iterrows():
            aid = arow['商品ID']
            if aid not in anomaly_summary:
                anomaly_summary[aid] = {'codes': set(), 'lines': 0}
            anomaly_summary[aid]['codes'].add(arow['商家编码'])
            anomaly_summary[aid]['lines'] += 1

        for rec in summary_records:
            pid = rec['_temp_id']
            if pid in anomaly_summary:
                rec['异常编码数量'] = len(anomaly_summary[pid]['codes'])
                rec['异常订单行数'] = anomaly_summary[pid]['lines']
            else:
                rec['异常编码数量'] = 0
                rec['异常订单行数'] = 0
            del rec['_temp_id']
    else:
        for rec in summary_records:
            rec['异常编码数量'] = 0
            rec['异常订单行数'] = 0
            del rec['_temp_id']

    # 补充：只存在于异常中的商品ID也加入汇总
    if len(anomaly_df) > 0:
        for aid, info in anomaly_summary.items():
            if aid not in normal_ids:
                summary_records.append({
                    '商品ID': aid,
                    '订单行数': 0,
                    '商品数量合计': 0,
                    '总收入': 0,
                    '总成本': 0,
                    '总利润': 0,
                    '异常编码数量': len(info['codes']),
                    '异常订单行数': info['lines'],
                })

    return detail_records, summary_records, anomaly_rows


# ============================================================
# Excel 输出
# ============================================================

def auto_width(ws, min_width=8, max_width=40):
    """自动调整列宽"""
    for col_idx in range(1, ws.max_column + 1):
        max_len = 0
        col_letter = get_column_letter(col_idx)
        for row_idx in range(1, ws.max_row + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            if cell.value:
                # 估算中文字符宽度（中文约占2个字符宽度）
                val = str(cell.value)
                length = 0
                for ch in val:
                    if '一' <= ch <= '鿿' or '　' <= ch <= '〿' or '＀' <= ch <= '￯':
                        length += 2
                    else:
                        length += 1
                max_len = max(max_len, length)
        width = min(max(max_len + 2, min_width), max_width)
        ws.column_dimensions[col_letter].width = width


def write_header(ws, headers, fill=None):
    """写表头并设置样式"""
    header_font = Font(bold=True, size=11)
    if fill is None:
        fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_font = Font(bold=True, size=11, color="FFFFFF")

    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = fill
        cell.alignment = Alignment(horizontal='center', vertical='center')


def write_sheet1_detail(ws, detail_records):
    """Sheet1: 商品ID编码明细"""
    # 检查是否有「来源文件」字段（多订单表合并时会有）
    has_source = any('来源文件' in rec for rec in detail_records)
    headers = [
        '商品ID', '商家编码', '成本单价', '编码对应订单行数',
        '编码对应商品数量', '编码对应总成本', '编码对应总收入',
        '编码对应总利润', 'ID对应总利润', '状态',
    ]
    if has_source:
        headers.append('来源文件')
    write_header(ws, headers)
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions

    for i, rec in enumerate(detail_records, 2):
        ws.cell(row=i, column=1, value=rec['商品ID'])
        ws.cell(row=i, column=2, value=rec['商家编码'])
        ws.cell(row=i, column=3, value=rec['成本单价'])
        ws.cell(row=i, column=4, value=rec['编码对应订单行数'])
        ws.cell(row=i, column=5, value=rec['编码对应商品数量'])
        c6 = ws.cell(row=i, column=6, value=rec['编码对应总成本'])
        c6.number_format = '#,##0.00'
        c7 = ws.cell(row=i, column=7, value=rec['编码对应总收入'])
        c7.number_format = '#,##0.00'
        c8 = ws.cell(row=i, column=8, value=rec['编码对应总利润'])
        c8.number_format = '#,##0.00'
        c9 = ws.cell(row=i, column=9, value=rec['ID对应总利润'])
        c9.number_format = '#,##0.00'
        ws.cell(row=i, column=10, value=rec['状态'])
        if has_source:
            ws.cell(row=i, column=11, value=rec.get('来源文件', ''))

    auto_width(ws)


def write_sheet2_summary(ws, summary_records):
    """Sheet2: 商品ID汇总"""
    headers = [
        '商品ID', '订单行数', '商品数量合计', '总收入',
        '总成本', '总利润', '异常编码数量', '异常订单行数',
    ]
    write_header(ws, headers)
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions

    for i, rec in enumerate(summary_records, 2):
        ws.cell(row=i, column=1, value=rec['商品ID'])
        ws.cell(row=i, column=2, value=rec['订单行数'])
        ws.cell(row=i, column=3, value=rec['商品数量合计'])
        c4 = ws.cell(row=i, column=4, value=rec['总收入'])
        c4.number_format = '#,##0.00'
        c5 = ws.cell(row=i, column=5, value=rec['总成本'])
        c5.number_format = '#,##0.00'
        c6 = ws.cell(row=i, column=6, value=rec['总利润'])
        c6.number_format = '#,##0.00'
        ws.cell(row=i, column=7, value=rec['异常编码数量'])
        ws.cell(row=i, column=8, value=rec['异常订单行数'])

    auto_width(ws)


def write_sheet3_anomaly(ws, anomaly_rows):
    """Sheet3: 异常明细"""
    has_source = any('来源文件' in rec for rec in anomaly_rows)
    headers = [
        '异常类型', '商品ID', '商家编码', '商品数量',
        '商家收入', '订单提交时间', '订单号', '原因',
    ]
    if has_source:
        headers.append('来源文件')
    # 异常明细用醒目样式
    header_fill = PatternFill(start_color="C00000", end_color="C00000", fill_type="solid")
    write_header(ws, headers, fill=header_fill)
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions

    for i, rec in enumerate(anomaly_rows, 2):
        ws.cell(row=i, column=1, value=rec['异常类型'])
        ws.cell(row=i, column=2, value=rec['商品ID'])
        ws.cell(row=i, column=3, value=rec['商家编码'])
        ws.cell(row=i, column=4, value=rec['商品数量'])
        ws.cell(row=i, column=5, value=rec['商家收入'])
        ws.cell(row=i, column=6, value=str(rec['订单提交时间']) if rec['订单提交时间'] else '')
        ws.cell(row=i, column=7, value=rec['订单号'])
        ws.cell(row=i, column=8, value=rec['原因'])
        if has_source:
            ws.cell(row=i, column=9, value=rec.get('来源文件', ''))

    auto_width(ws)


def write_sheet4_cost_snapshot(ws, cost_records):
    """Sheet4: 成本表快照"""
    headers = ['商家编码', '成本单价', '状态', '原因']
    write_header(ws, headers)
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions

    for i, rec in enumerate(cost_records, 2):
        ws.cell(row=i, column=1, value=rec['商家编码'])
        c2 = ws.cell(row=i, column=2, value=rec['成本单价'])
        if rec['成本单价'] is not None:
            c2.number_format = '#,##0.00'
        ws.cell(row=i, column=3, value=rec['状态'])
        ws.cell(row=i, column=4, value=rec['原因'])

    auto_width(ws)


def generate_output_excel(detail_records, summary_records, anomaly_rows,
                          cost_snapshot_records, output_filepath,
                          per_file_results=None):
    """生成输出 Excel 文件

    参数:
        per_file_results: 可选，每个订单表的独立结果列表
                          [(store_name, detail, summary, anomaly), ...]
                          多订单表时按店铺分 sheet，单文件时为 None
    """
    wb = Workbook()

    # 删除默认 sheet
    wb.remove(wb.active)

    # 多订单表时：每个店铺独立的明细/汇总/异常 + 总计 sheet
    if per_file_results and len(per_file_results) > 1:
        used_names = set()

        for store_name, detail, summary, anomaly in per_file_results:
            # 提取店铺简称（去扩展名）
            base_name = os.path.splitext(store_name)[0]
            # 处理重名
            unique_name = base_name
            counter = 1
            while unique_name in used_names:
                unique_name = f"{base_name}_{counter}"
                counter += 1
            used_names.add(unique_name)

            # 确保完整 sheet 名不超过 31 字符
            max_prefix_len = 31 - len('-明细')
            if len(unique_name) > max_prefix_len:
                unique_name = unique_name[:max_prefix_len]

            # 店铺明细
            ws_detail = wb.create_sheet(f'{unique_name}-明细')
            write_sheet1_detail(ws_detail, detail)

            # 店铺汇总
            ws_summary = wb.create_sheet(f'{unique_name}-汇总')
            write_sheet2_summary(ws_summary, summary)

            # 店铺异常
            ws_anomaly = wb.create_sheet(f'{unique_name}-异常')
            write_sheet3_anomaly(ws_anomaly, anomaly)

        # 总计 sheet（合并所有订单表）
        ws_total_detail = wb.create_sheet('总计-明细')
        write_sheet1_detail(ws_total_detail, detail_records)

        ws_total_summary = wb.create_sheet('总计-汇总')
        write_sheet2_summary(ws_total_summary, summary_records)

        ws_total_anomaly = wb.create_sheet('总计-异常')
        write_sheet3_anomaly(ws_total_anomaly, anomaly_rows)

        # 成本表快照
        ws_cost = wb.create_sheet('成本表快照')
        write_sheet4_cost_snapshot(ws_cost, cost_snapshot_records)
    else:
        # 单文件时保持原有4-sheet结构
        ws1 = wb.create_sheet('商品ID编码明细')
        write_sheet1_detail(ws1, detail_records)

        ws2 = wb.create_sheet('商品ID汇总')
        write_sheet2_summary(ws2, summary_records)

        ws3 = wb.create_sheet('异常明细')
        write_sheet3_anomaly(ws3, anomaly_rows)

        ws4 = wb.create_sheet('成本表快照')
        write_sheet4_cost_snapshot(ws4, cost_snapshot_records)

    wb.save(output_filepath)
    return output_filepath


# ============================================================
# 主流程
# ============================================================

def run_calculation(order_filepaths, cost_filepath, output_dir,
                    start_date=None, end_date=None, order_statuses=None,
                    log_func=None, progress_func=None):
    """执行完整的利润计算流程（支持多个订单表 + 同一成本表）

    参数:
        order_filepaths: 订单表路径列表 (list of str)，支持多个订单表
        cost_filepath: 成本表路径
        output_dir: 输出目录
        start_date: 开始日期 (datetime 或 None)
        end_date: 结束日期 (datetime 或 None)
        order_statuses: 订单状态列表 (list of str 或 None)，支持多选
        log_func: 日志回调函数
        progress_func: 进度回调函数，签名 (percent: int, message: str)

    返回:
        (output_path, stats_dict)
    """
    def log(msg):
        if log_func:
            log_func(msg)
        else:
            print(msg)

    def progress(percent, message):
        if progress_func:
            progress_func(percent, message)

    log("=" * 60)
    log("开始利润计算...")
    progress(2, "正在读取成本表...")

    # 1. 处理成本表（所有订单表共用）
    log(f"\n[1/5] 读取成本表: {os.path.basename(cost_filepath)}")
    cost_map, cost_snapshot = process_cost_table(cost_filepath)
    log(f"  成本表行数: {len(cost_snapshot)}")
    normal_cost = sum(1 for r in cost_snapshot if r['状态'] == '正常')
    dup_cost = sum(1 for r in cost_snapshot if r['状态'] == '成本重复')
    empty_cost = sum(1 for r in cost_snapshot if r['状态'] == '编码为空')
    bad_cost = sum(1 for r in cost_snapshot if r['状态'] == '成本异常')
    log(f"  正常: {normal_cost}, 成本重复: {dup_cost}, 编码为空: {empty_cost}, 成本异常: {bad_cost}")
    log(f"  可用于匹配的编码数量: {len(cost_map)}")
    progress(8, "成本表处理完成")

    # 2. 遍历处理每个订单表
    # 进度区间: 10% ~ 80%，按订单表数量均分
    ORDER_PROGRESS_START = 10
    ORDER_PROGRESS_END = 80
    total_order_rows = 0
    all_detail_records = []
    all_summary_records = []
    all_anomaly_rows = []
    per_file_results = []  # 每个订单表的独立结果: [(store_name, detail, summary, anomaly), ...]

    num_orders = len(order_filepaths)
    for idx, order_filepath in enumerate(order_filepaths):
        order_name = os.path.basename(order_filepath)
        order_base = ORDER_PROGRESS_START + (ORDER_PROGRESS_END - ORDER_PROGRESS_START) * idx // num_orders
        order_next_base = ORDER_PROGRESS_START + (ORDER_PROGRESS_END - ORDER_PROGRESS_START) * (idx + 1) // num_orders

        log(f"\n[2.{idx+1}/5] 读取订单表 [{idx+1}/{num_orders}]: {order_name}")

        # 读取阶段（最耗时，约占订单表处理时间的 90%）
        progress(order_base + 3, f"正在读取订单表 ({order_name})，文件较大请耐心等待...")
        df_order = pd.read_excel(order_filepath)
        total_order_rows += len(df_order)
        log(f"  订单表行数: {len(df_order)}")
        progress(order_base + 12, f"订单表读取完成 ({order_name})，正在识别字段...")

        # 识别字段（只识别一次，传给 process_orders 避免重复识别）
        col_map = identify_order_columns(df_order)
        if idx == 0:
            log("  识别字段（首个订单表）:")
            for key, val in col_map.items():
                log(f"    {key} -> [{val}]")

        # 执行利润计算（传入已读取的 df 和 col_map，避免重复 read_excel + 重复识别）
        progress(order_base + 14, f"正在处理订单数据 ({order_name})...")
        detail_records, summary_records, anomaly_rows = process_orders(
            order_filepath, cost_map, start_date, end_date, order_statuses,
            df=df_order, col_map=col_map
        )
        log(f"  正常编码明细节: {len(detail_records)} 行")
        log(f"  商品ID汇总: {len(summary_records)} 行")
        log(f"  异常明细: {len(anomaly_rows)} 行")

        # 保存每个订单表的独立结果（用于按店铺分sheet输出）
        source_name = os.path.basename(order_filepath)
        per_file_results.append((
            source_name,
            [dict(r) for r in detail_records],      # 浅拷贝，不含「来源文件」
            list(summary_records),
            [dict(r) for r in anomaly_rows],         # 浅拷贝，不含「来源文件」
        ))

        # 给明细记录标记来源文件（用于合并后总计sheet追溯）
        for rec in detail_records:
            rec['来源文件'] = source_name
        for rec in anomaly_rows:
            rec['来源文件'] = source_name

        all_detail_records.extend(detail_records)
        all_summary_records.extend(summary_records)
        all_anomaly_rows.extend(anomaly_rows)

        progress(order_next_base, f"订单表 {order_name} 处理完成")

    # 3. 合并汇总：同商品ID的汇总行需要合并
    log(f"\n[3/5] 合并多个订单表数据...")
    log(f"  合并前: 明细节 {len(all_detail_records)} 行, 汇总 {len(all_summary_records)} 行, 异常 {len(all_anomaly_rows)} 行")
    progress(82, "正在合并多个订单表数据...")

    # 合并 detail records：相同 (商品ID, 商家编码) 的需要重新聚合
    if len(order_filepaths) > 1 and len(all_detail_records) > 0:
        detail_df = pd.DataFrame(all_detail_records)
        grouped = detail_df.groupby(['商品ID', '商家编码'], dropna=False)

        merged_detail = []
        for (pid, mcode), grp in grouped:
            order_lines = grp['编码对应订单行数'].sum()
            total_qty = grp['编码对应商品数量'].sum()
            cost_price = grp['成本单价'].iloc[0]
            total_cost = cost_price * total_qty
            total_income = grp['编码对应总收入'].sum()
            total_profit = total_income - total_cost
            sources = ', '.join(sorted(set(grp['来源文件'].dropna())))

            merged_detail.append({
                '商品ID': pid,
                '商家编码': mcode,
                '成本单价': cost_price,
                '编码对应订单行数': int(order_lines),
                '编码对应商品数量': total_qty,
                '编码对应总成本': round(total_cost, 2),
                '编码对应总收入': round(total_income, 2),
                '编码对应总利润': round(total_profit, 2),
                '_total_profit': total_profit,
                '来源文件': sources,
                '状态': '正常',
            })

        # 计算 ID对应总利润
        id_profit_map = {}
        for rec in merged_detail:
            pid = rec['商品ID']
            id_profit_map[pid] = id_profit_map.get(pid, 0) + rec['_total_profit']

        for rec in merged_detail:
            rec['ID对应总利润'] = round(id_profit_map[rec['商品ID']], 2)
            del rec['_total_profit']

        all_detail_records = merged_detail

    # 合并 summary records：相同 商品ID 的需要合并
    if len(order_filepaths) > 1 and len(all_summary_records) > 0:
        summary_df = pd.DataFrame(all_summary_records)
        grouped = summary_df.groupby('商品ID', dropna=False)

        merged_summary = []
        for pid, grp in grouped:
            merged_summary.append({
                '商品ID': pid,
                '订单行数': int(grp['订单行数'].sum()),
                '商品数量合计': grp['商品数量合计'].sum(),
                '总收入': round(grp['总收入'].sum(), 2),
                '总成本': round(grp['总成本'].sum(), 2),
                '总利润': round(grp['总利润'].sum(), 2),
                '异常编码数量': int(grp['异常编码数量'].max()),  # 取最大，因为不同文件可能有相同异常编码
                '异常订单行数': int(grp['异常订单行数'].sum()),
            })
        all_summary_records = merged_summary

    log(f"  合并后: 明细节 {len(all_detail_records)} 行, 汇总 {len(all_summary_records)} 行, 异常 {len(all_anomaly_rows)} 行")

    # 4. 输出
    log("\n[4/5] 生成输出 Excel...")
    progress(88, "正在生成输出 Excel 文件...")

    # 生成文件名
    if start_date and end_date:
        start_str = start_date.strftime('%Y-%m-%d')
        end_str = end_date.strftime('%Y-%m-%d')
        filename = f"利润统计结果_{start_str}_{end_str}"
    else:
        filename = "利润统计结果_全部数据"

    if order_statuses:
        status_str = '_'.join(order_statuses)
        # 文件名不能太长
        if len(status_str) > 50:
            status_str = status_str[:47] + '...'
        filename += f"_{status_str}"

    # 多订单表时追加标识
    if len(order_filepaths) > 1:
        filename += f"_共{len(order_filepaths)}个订单表"

    filename += ".xlsx"

    os.makedirs(output_dir, exist_ok=True)
    output_path = os.path.join(output_dir, filename)
    generate_output_excel(all_detail_records, all_summary_records, all_anomaly_rows,
                          cost_snapshot, output_path,
                          per_file_results=per_file_results)

    progress(96, "输出文件生成完成")

    log(f"\n{'=' * 60}")
    log(f"计算完成!")
    log(f"输出文件: {output_path}")

    normal_lines = sum(r['编码对应订单行数'] for r in all_detail_records)

    stats = {
        '订单表行数': total_order_rows,
        '成本表行数': len(cost_snapshot),
        '正常参与计算行数': normal_lines,
        '异常行数': len(all_anomaly_rows),
        '输出文件': output_path,
    }

    log(f"\n统计摘要:")
    for k, v in stats.items():
        if k != '输出文件':
            log(f"  {k}: {v}")
    log(f"  输出文件: {output_path}")

    progress(100, "计算完成!")
    return output_path, stats


# ============================================================
# 命令行测试入口
# ============================================================

def parse_date(s):
    """解析日期字符串 YYYY-MM-DD，返回 datetime"""
    if not s or not s.strip():
        return None
    s = s.strip()
    try:
        return datetime.strptime(s, '%Y-%m-%d')
    except ValueError:
        raise ValueError(f"日期格式错误: {s}，应为 YYYY-MM-DD")


def main():
    """命令行测试入口"""
    # 默认路径
    base_dir = os.path.dirname(os.path.abspath(__file__))
    order_file = os.path.join(base_dir, '咔咔唛.xlsx')
    cost_file = os.path.join(base_dir, '24日成本表.xlsx')
    output_dir = os.path.join(base_dir, 'output')

    # 检查文件存在
    if not os.path.exists(order_file):
        print(f"错误: 订单表不存在: {order_file}")
        return 1
    if not os.path.exists(cost_file):
        print(f"错误: 成本表不存在: {cost_file}")
        return 1

    start_date = None
    end_date = None

    # 支持命令行日期参数: python profit_core.py [start_date] [end_date]
    import sys
    if len(sys.argv) >= 2:
        start_date = parse_date(sys.argv[1])
    if len(sys.argv) >= 3:
        end_date = parse_date(sys.argv[2])
        # 结束日期自动设为当天 23:59:59
        end_date = end_date.replace(hour=23, minute=59, second=59)

    try:
        output_path, stats = run_calculation(
            [order_file], cost_file, output_dir,
            start_date=start_date, end_date=end_date
        )
        print("\n" + "=" * 60)
        print("测试通过")
        print(f"订单表行数：{stats['订单表行数']}")
        print(f"成本表行数：{stats['成本表行数']}")
        print(f"正常参与计算行数：{stats['正常参与计算行数']}")
        print(f"异常行数：{stats['异常行数']}")
        print(f"输出文件：{stats['输出文件']}")
        return 0
    except Exception as e:
        print(f"\n测试失败: {e}")
        import traceback
        traceback.print_exc()
        return 1


if __name__ == '__main__':
    exit(main())
